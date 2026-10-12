#!/usr/bin/env python3
"""
将 Git 提交历史导出为格式优化过的 Excel 文件（默认忽略合并提交，文件名自动包含日期）。
支持在仓库的任意子目录下运行，也可指定仓库路径。

依赖：
    pip install gitpython openpyxl pandas

用法：
    python gitlog_to_excel.py [--repo PATH] [--output OUTPUT] [--branch BRANCH] [--max-count N] [--since DATE] [--author AUTHOR] [--include-merges]

输出列：
    author        : 作者
    author_email  : 作者邮箱
    date          : 提交日期
    message       : 提交信息
    files         : 修改文件数量
    commit_hash   : 完整哈希
"""

import argparse
import os
import pandas as pd
from git import Repo, InvalidGitRepositoryError, NoSuchPathError
from datetime import datetime
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter


def parse_args():
    parser = argparse.ArgumentParser(description="导出 Git log 到格式优化的 Excel（默认忽略合并提交）")
    parser.add_argument("--repo", "-r", default=".",
                        help="Git 仓库路径（默认当前目录，可在子目录运行）")
    parser.add_argument("--output", "-o", default=None,
                        help="输出 Excel 文件名（默认：git_log_YYYY-MM-DD.xlsx）")
    parser.add_argument("--branch", "-b", default="HEAD", help="要导出的分支或引用 (默认: HEAD)")
    parser.add_argument("--max-count", "-n", type=int, help="最多导出的提交数量")
    parser.add_argument("--since", help="起始日期，例如 '2024-01-01' 或 '2 weeks ago'")
    parser.add_argument("--author", help="按作者筛选")
    parser.add_argument("--grep", help="按提交信息关键字筛选")
    parser.add_argument("--include-merges", action="store_true",
                        help="包含合并提交（默认忽略）")
    return parser.parse_args()


def auto_format_excel(filepath):
    """优化 Excel 格式：自动列宽、表头加粗居中、添加边框等。"""
    wb = load_workbook(filepath)
    ws = wb.active

    # 表头样式
    header_font = Font(bold=True, size=11)
    header_alignment = Alignment(horizontal="center", vertical="center")
    header_fill = PatternFill(start_color="D9E2F3", end_color="D9E2F3", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )
    data_alignment = Alignment(vertical="center")

    # 设置表头样式
    for col_idx in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.alignment = header_alignment
        cell.fill = header_fill
        cell.border = thin_border

    # 设置数据区域样式并自动调整列宽
    for col_idx in range(1, ws.max_column + 1):
        max_length = 0
        column_letter = get_column_letter(col_idx)
        for row_idx in range(1, ws.max_row + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = thin_border
            if row_idx > 1:
                cell.alignment = data_alignment
                try:
                    cell_len = len(str(cell.value))
                    if cell_len > max_length:
                        max_length = cell_len
                except:
                    pass
        adjusted_width = min(max(max_length + 4, 12), 50)
        ws.column_dimensions[column_letter].width = adjusted_width

    ws.freeze_panes = "A2"
    wb.save(filepath)


def main():
    args = parse_args()

    # 确定输出文件名
    if args.output is None:
        today = datetime.now().strftime("%Y-%m-%d")
        output_file = f"git_log_{today}.xlsx"
    else:
        output_file = args.output

    # 定位 Git 仓库（支持子目录向上查找）
    try:
        repo = Repo(args.repo, search_parent_directories=True)
    except (InvalidGitRepositoryError, NoSuchPathError) as e:
        print(f"错误：未找到 Git 仓库（路径：{args.repo}）。请确保在仓库内或指定正确的仓库路径。")
        return

    print(f"使用仓库：{repo.working_dir}")

    kwargs = {}
    if args.max_count:
        kwargs["max_count"] = args.max_count
    if args.since:
        kwargs["since"] = args.since
    if args.author:
        kwargs["author"] = args.author
    if args.grep:
        kwargs["grep"] = args.grep

    commits = list(repo.iter_commits(args.branch, **kwargs))
    total = len(commits)
    print(f"找到 {total} 条提交记录（包含合并提交）")

    rows = []
    skipped = 0
    for commit in commits:
        if not args.include_merges and len(commit.parents) > 1:
            skipped += 1
            continue

        author = commit.author.name
        author_email = commit.author.email
        date = datetime.fromtimestamp(commit.committed_date).strftime("%Y-%m-%d %H:%M:%S")
        message = commit.summary
        commit_hash = commit.hexsha

        try:
            files_changed = commit.stats.total["files"]
        except Exception:
            files_changed = 0

        rows.append({
            "author": author,
            "author_email": author_email,
            "date": date,
            "message": message,
            "files": files_changed,
            "commit_hash": commit_hash,
        })

    if skipped:
        print(f"已忽略 {skipped} 条合并提交")

    df = pd.DataFrame(rows)
    column_order = ["author", "author_email", "date", "message", "files", "commit_hash"]
    df = df[column_order]

    df.to_excel(output_file, index=False, engine="openpyxl")
    print(f"已导出 {len(df)} 条普通提交到 {output_file}")

    auto_format_excel(output_file)
    print("Excel 格式已优化（自适应列宽、表头加粗、边框、冻结首行等）。")


if __name__ == "__main__":
    main()